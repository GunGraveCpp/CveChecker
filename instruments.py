import vulners
import os
import sys
import functools
import requests
from docx.shared import Cm
from bs4 import BeautifulSoup
from googletrans import Translator

from config import *


def time_measurement(func):
    '''Декоратор для замера времени'''

    import time
    @functools.wraps(func)
    def inner(*args, **kwargs):
        start = time.time()
        ret_value = func(*args, **kwargs)
        end = time.time()
        print('Function {} worked: {} min {} sec'.format(func.__name__,
            int(end - start) // 60, int(end - start) % 60))
        return ret_value
    return inner if DEBUG else func


def trace(func):
    '''Декоратор для отладки'''

    @functools.wraps(func)
    def inner(*args, **kwargs):
        ret_value = func(*args, **kwargs)
        print('Function {} \n\tInput args: {} \n\tInput kwargs: {} \n\tReturn {}'.format(
            func.__name__, args, kwargs, ret_value), end='\n\n')
        return ret_value
    return inner if DEBUG else func



def read_package_in_file(name):
    package_list = {'deb': [], 'rpm': [], 'archive': [], 'unknown': []}

    try:
        with open(name, 'r') as file:
            for package in file:
                if package.rfind('.deb') != -1:
                    package_list['deb'].append(
                        package.replace('_', ' ').replace('.deb', ''))
                elif package.rfind('.rpm') != -1:
                    package_list['rpm'].append(package.replace('.rpm', ''))
                # package.count('-') костыль, чтобы была возможность делить имя и версию архива
                elif (package.rfind('.tar') & package.rfind('.zip') != -1) & package.count('-'):
                    package_list['archive'].append(
                        package[:package.find('.tar') & package.find('.zip')])
                else:
                    package_list['unknown'].append(package)
    except FileNotFoundError as er:
        print(er, file=sys.stderr)

    return package_list


def set_col_widths(table):
    widths = (Cm(width) for width in COLUMN_SIZE)
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def generate_doc_table(document):
    document.add_heading(TABLE_NAME, 1)
    table = document.add_table(rows=1, cols=5, style='Table Grid')

    for column in range(0, len(table.columns)):
        table.cell(0, column).text = TABLE_HEADER[column]

    return table


def translate_to_rus(description):
    translator = Translator()
    translated = translator.translate(description, dest='ru')
    return translated.text


def deb_version(name):
    return name.split()[1]


def rpm_version(name):
    return name.rsplit('.', 1)[0].rsplit('-', 1)[0].rsplit('-', 1)[1]


def archive_version(name):
    return name.split('-', 1)[1]


def get_description(cve):
    '''  Функция получения описпния CVE

    Получение описания и наличия eploit с https://nvd.nist.gov/
    по идентификатору cve

    '''

    req = requests.get(NIST_LINK.format(cve))
    answ = {'description': WITHOUT_DESCRIPTION,
            'exploit': AVAILABILITY_EPLOIT_NO}

    if req.ok:
        soup = BeautifulSoup(req.text, 'lxml')
        span_list = soup.find_all('span', class_="badge")
        if span_list:
            for span in span_list:
                if span.text == 'Exploit':
                    answ['exploit'] = AVAILABILITY_EPLOIT_YES

        div = soup.find('div', id=DIV_ID)
        if div:
            desc = div.find('p', {'data-testid': 'vuln-description'})
            if desc:
                answ['description'] = desc.text

    return answ


def fill_package_info(package_info, cve_list):
    '''Заполнение информации о cve'''

    for cve in cve_list:
        cve_info = get_description(cve)

        package_info['cvelist'].append(
            {'cve': cve,
             'description': translate_to_rus(cve_info['description']),
             'exploit': cve_info['exploit']
             }
        )


def get_package_info_deb(name, vulners_api):
    '''Получение информации о deb пакете'''

    package_info = {'package': name.replace(
        ' ', '_'), 'version': deb_version(name), 'cvelist': []}
    results = vulners_api.audit(
        os=DEB_OS, os_version=DEB_VERSION, package=[name])

    if len(results['reasons']) != 0:
        cve_list = results.get('cvelist')
        fill_package_info(package_info, cve_list)

    return package_info


def get_package_info_rpm(name, vulners_api):
    '''Получение информации о rpm пакете'''

    package_info = {'package': name,
                    'version': rpm_version(name), 'cvelist': []}

    results = vulners_api.audit(
        os=CENT_OS, os_version=CENT_VERSION, package=[name])

    if len(results['reasons']) != 0:
        cve_list = results.get('cvelist')
        fill_package_info(package_info, cve_list)

    return package_info


def get_package_info_archive(name, vulners_api):
    '''Получение информации об архиве'''

    package_info = {'package': name.split(
        '-', 1)[0], 'version': archive_version(name), 'cvelist': []}

    results = vulners_api.softwareVulnerabilities(
        package_info['package'].upper(), package_info['version'])

    if len(results) != 0:
        cve_list = [cve['id'].rsplit(':', 1)[1] for cve in results['software']]
        fill_package_info(package_info, cve_list)

    return package_info


def past_package_in_table(table, package_info):
    '''Вставки информации о пакете в таблицу'''
    idx_first_row = len(table.rows)
    if len(package_info['cvelist']) != 0:
        for cve_info in package_info['cvelist']:
            row = table.add_row()
            row.cells[2].text = cve_info['cve']
            row.cells[3].text = cve_info['description']
            row.cells[4].text = cve_info['exploit']

        table.cell(idx_first_row, 0).merge(table.cell(len(table.rows) - 1, 0))
        table.cell(idx_first_row, 1).merge(table.cell(len(table.rows) - 1, 1))
    else:
        row = table.add_row()
        row.cells[2].text = INVULNERABILITY_MESSAGE
        row.cells[3].text = INVULNERABILITY_MESSAGE
        row.cells[4].text = INVULNERABILITY_MESSAGE

    table.cell(idx_first_row, 0).text = package_info['package']
    table.cell(idx_first_row, 1).text = package_info['version']


@time_measurement
def run_check(table, package_list):
    '''Запуск задачи проверки и генерация отчета'''

    vulners_api = vulners.Vulners(api_key=os.environ['VULNERS_API_KEY'])

    for package_deb in package_list['deb']:
        package_info = get_package_info_deb(package_deb, vulners_api)
        past_package_in_table(table, package_info)
    for package_rpm in package_list['rpm']:
        package_info = get_package_info_rpm(package_rpm, vulners_api)
        past_package_in_table(table, package_info)
    for package_archive in package_list['archive']:
        try:
            package_info = get_package_info_archive(package_archive, vulners_api)
            past_package_in_table(table, package_info)
        except Exception as error:
            print(error)
    with open("unknown.txt", "w") as outfile:
        outfile.write('\n'.join(package_list['unknown']))
